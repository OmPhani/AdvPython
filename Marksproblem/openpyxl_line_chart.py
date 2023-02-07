import openpyxl
from openpyxl.styles import Border,Side,Font,PatternFill,Alignment
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
wb = Workbook()
sheet = wb.active
chart = LineChart()

sales_data = [
              ["Years","Sales"],
              [2010,10000],
              [2011,90000],
              [2012,20000],
              [2013,25000],
              [2014,44000]
]

for row in sales_data:
    sheet.append(row)

data = Reference(sheet, min_col=1, max_col=2, min_row=1, max_row=7)
chart.add_data(data, titles_from_data=True)
sheet.add_chart(chart,"E2")

wb.save("C://Users/user790/Desktop/linechartdata.xlsx")